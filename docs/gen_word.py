from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

def set_chinese_font(run, font_name='SimSun', size=10.5, bold=False):
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def add_heading_zh(doc, text, level=1):
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    font_name = 'SimHei' if level <= 2 else 'SimSun'
    set_chinese_font(run, font_name=font_name, size=(18 if level==1 else (14 if level==2 else 12)), bold=True)
    return heading

def add_paragraph_zh(doc, text, bold=False, size=10.5, alignment=None):
    p = doc.add_paragraph()
    if alignment:
        p.alignment = alignment
    run = p.add_run(text)
    set_chinese_font(run, bold=bold, size=size)
    return p

def add_screenshot_placeholder(doc, desc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'【截图占位：{desc}】')
    set_chinese_font(run, font_name='SimSun', size=10, bold=False)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('（请替换为实际页面截图）')
    set_chinese_font(run, font_name='SimSun', size=9, bold=False)
    run.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA0)

def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    style.font.size = Pt(10.5)

    # Cover
    add_paragraph_zh(doc, '', size=20)
    add_paragraph_zh(doc, '', size=20)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('数据运维数字员工平台')
    set_chinese_font(run, font_name='SimHei', size=26, bold=True)
    add_paragraph_zh(doc, '', size=10)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run('产品详细设计文档')
    set_chinese_font(run, font_name='SimHei', size=22, bold=True)
    add_paragraph_zh(doc, '', size=30)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run('版本：V1.0.0\n日期：2026-06-02')
    set_chinese_font(run, font_name='SimSun', size=12)
    doc.add_page_break()

    # Revision history
    add_heading_zh(doc, '文档修订记录', level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(['版本号', '修订日期', '修订内容', '作者', '审核人']):
        run = hdr[i].paragraphs[0].add_run(h)
        set_chinese_font(run, bold=True)
    row = table.add_row().cells
    for i, v in enumerate(['V1.0.0', '2026-06-02', '初稿完成', '-', '-']):
        run = row[i].paragraphs[0].add_run(v)
        set_chinese_font(run)
    doc.add_page_break()

    # Chapter 1
    add_heading_zh(doc, '1 产品概述', level=1)
    add_paragraph_zh(doc, '数据运维数字员工平台是一款面向数据运维场景的智能化管理平台，通过 AI 辅助实现数据质量稽核、根因分析、月账管理、系统配置等核心能力，降低人工操作成本，提升数据运维效率与质量。')
    add_paragraph_zh(doc, '产品定位：面向企业数据运维团队的智能化协作平台。')
    add_paragraph_zh(doc, '目标用户：数据运维工程师、数据质量管理员、月账处理人员、系统管理员。')

    add_heading_zh(doc, '1.1 功能架构', level=2)
    add_paragraph_zh(doc, '平台包含五大功能模块：')
    add_paragraph_zh(doc, '• 总览看板：提供全局数据视图与核心指标监控。')
    add_paragraph_zh(doc, '• 数据质量稽核：覆盖指标配置、AI 规则生成、规则确认、告警管理、结果视图。')
    add_paragraph_zh(doc, '• 根因分析：提供知识库、智能分析、任务列表、案例库、改进建议。')
    add_paragraph_zh(doc, '• 月账管理：覆盖出账日报、报表发布、月账进度、月账配置。')
    add_paragraph_zh(doc, '• 系统设置：工具注册、Prompt 管理、知识库、通知配置、用户管理。')
    doc.add_page_break()

    # Chapter 2
    add_heading_zh(doc, '2 系统架构', level=1)
    add_heading_zh(doc, '2.1 技术架构', level=2)
    add_paragraph_zh(doc, '前端：React 18 + TypeScript + Ant Design + Vite')
    add_paragraph_zh(doc, '后端：Python FastAPI + SQLAlchemy + Alembic + MySQL')
    add_paragraph_zh(doc, 'AI 能力：集成 OpenAI / Claude API，提供智能规则生成、根因分析、日报生成等能力')
    add_paragraph_zh(doc, '部署：Nginx 反向代理 + Uvicorn ASGI 服务')
    add_heading_zh(doc, '2.2 部署架构', level=2)
    add_paragraph_zh(doc, '• Web 服务：Nginx → Vite 前端静态资源 / FastAPI 后端 API')
    add_paragraph_zh(doc, '• 数据库：MySQL 8.0（字符集 utf8mb4）')
    add_paragraph_zh(doc, '• 缓存/队列：Redis（Celery Broker / Result Backend）')
    doc.add_page_break()

    # Chapter 3
    add_heading_zh(doc, '3 功能模块详细设计', level=1)

    modules = [
        ('3.1 总览看板', '/dashboard', '展示全局核心指标、任务运行状态、告警统计、快捷入口等。', [
            '顶部统计卡片：任务总数、运行中、已完成、异常数',
            '中间图表区：任务趋势图、异常分布图',
            '底部列表：最近告警、待办任务',
        ]),
        ('3.2 数据质量稽核', None, '', []),
        ('3.2.1 指标配置', '/audit/field-config', '管理数据质量稽核的字段配置，包括数据源、表、字段元信息维护。', [
            '字段列表：支持增删改查',
            '字段详情：字段类型、长度、可空性、样例数据',
            '批量导入/导出',
        ]),
        ('3.2.2 AI规则生成', '/audit/rule-generate', '基于字段元数据与样例数据，通过 AI 自动生成稽核规则。', [
            '选择目标字段',
            'AI 生成规则建议（空值检查、重复检查、范围检查、格式检查等）',
            '规则预览与编辑',
            '一键保存到规则库',
        ]),
        ('3.2.3 规则确认', '/audit/rule-confirm', '对 AI 生成的规则进行人工确认或驳回，确保规则准确性。', [
            '待确认规则列表',
            '规则详情对比（AI 生成 vs 人工调整）',
            '批量确认/驳回',
        ]),
        ('3.2.4 告警管理', '/audit/alert-manage', '管理数据质量稽核产生的告警，支持升级规则配置。', [
            '告警列表：按级别、状态、类型筛选',
            '告警详情：关联规则、异常数据、处理记录',
            '告警升级规则配置',
        ]),
        ('3.2.5 结果视图', '/audit/result-view', '查看稽核执行结果与统计报表。', [
            '执行记录列表',
            '异常数据明细',
            '通过率统计',
            '报告导出',
        ]),
        ('3.3 根因分析', None, '', []),
        ('3.3.1 知识库', '/root-cause/knowledge-base', '管理运维知识文档与常见问题解决方案。', [
            '知识分类与检索',
            '文档上传与在线编辑',
            '版本管理',
        ]),
        ('3.3.2 智能分析', '/root-cause/analysis', '基于 AI 对问题进行智能根因分析。', [
            '输入问题描述',
            'AI 自动匹配知识库案例',
            '生成分析路径与结论',
            '保存为案例',
        ]),
        ('3.3.3 任务列表', '/root-cause/task-list', '管理根因分析任务。', [
            '任务列表与状态',
            '任务指派与跟踪',
            '分析结果归档',
        ]),
        ('3.3.4 案例库', '/root-cause/case-library', '沉淀历史问题案例，支持检索与复用。', [
            '案例分类与标签',
            '案例检索与推荐',
            '使用统计与评分',
        ]),
        ('3.3.5 改进建议', '/root-cause/suggestion', '基于历史案例与趋势生成改进建议。', [
            '建议列表',
            '建议采纳跟踪',
            '效果评估',
        ]),
        ('3.4 月账管理', None, '', []),
        ('3.4.1 出账日报', '/monthly/daily-report', '生成与查看每日出账报告。', [
            '日报列表',
            '日报详情（任务完成情况、异常统计）',
            'AI 辅助生成日报',
            '日报发布',
        ]),
        ('3.4.2 报表发布', '/monthly/report-publish', '管理月总结报表的生成与发布。', [
            '报表模板管理',
            '报表生成与预览',
            '发布审批',
            '发布历史',
        ]),
        ('3.4.3 月账进度', '/monthly/ledger-display', '展示月账作业的四级层次结构（阶段→里程碑→作业计划→任务）及整体进度。', [
            '账期选择：支持从 202512 到当前账期的倒排选择',
            '总体进度圆环图',
            '作业阶段卡片（可点击下钻）',
            '里程碑列表（展开显示作业计划与任务）',
            '任务类型标签（TDP任务、发布消息、人工操作、SQL脚本）',
        ]),
        ('3.4.4 月账配置', '/monthly/config', '管理月账作业的阶段、里程碑、作业计划与任务配置。', [
            '阶段配置',
            '里程碑配置',
            '作业计划配置',
            '任务配置',
            '模板解析（从自然语言模板自动解析为结构化任务）',
        ]),
        ('3.5 系统设置', None, '', []),
        ('3.5.1 工具注册', '/settings/tools', '注册与管理外部工具接口。', [
            '工具列表',
            '工具参数配置',
            '调用日志',
        ]),
        ('3.5.2 Prompt管理', '/settings/prompts', '管理 AI 使用的 Prompt 模板。', [
            'Prompt 列表',
            '版本管理',
            '效果测试',
        ]),
        ('3.5.3 知识库', '/settings/knowledge-base', '系统级知识库管理。', [
            '知识分类',
            '文档管理',
            '向量检索配置',
        ]),
        ('3.5.4 通知配置', '/settings/notifications', '配置通知渠道与模板。', [
            '通知渠道（系统/邮件/短信/钉钉/企业微信）',
            '通知模板',
            '订阅规则',
        ]),
        ('3.5.5 用户管理', '/settings/user-manage', '管理系统用户、角色与权限。', [
            '用户列表',
            '角色管理',
            '权限配置',
            '部门管理',
        ]),
    ]

    for title, path, desc, features in modules:
        if path is None:
            add_heading_zh(doc, title, level=2 if title.startswith('3.') and '.' in title.split(' ')[0] else 1)
            continue
        level = 2 if title.startswith('3.') and len(title.split(' ')[0].split('.')) <= 2 else 3
        add_heading_zh(doc, title, level=level)
        if desc:
            add_paragraph_zh(doc, desc)
        add_paragraph_zh(doc, '页面路径：' + (path or '-'))
        if features:
            add_paragraph_zh(doc, '功能要点：')
            for f in features:
                add_paragraph_zh(doc, '• ' + f)
        add_screenshot_placeholder(doc, title.split(' ', 1)[1] + '页面')
        doc.add_paragraph()

    doc.add_page_break()

    # Chapter 4
    add_heading_zh(doc, '4 数据库设计', level=1)
    add_paragraph_zh(doc, '数据库采用 MySQL 8.0，字符集 utf8mb4，共涉及 5 个业务域、约 30+ 张表。')
    add_heading_zh(doc, '4.1 表清单', level=2)
    tables = [
        ('数据质量稽核', [
            'dq_audit_field_config', 'dq_audit_rule_config', 'dq_audit_task_config',
            'dq_task_importance_config', 'dq_alert_upgrade_rule', 'dq_audit_execution',
            'dq_audit_exception', 'dq_audit_report',
        ]),
        ('根因分析', [
            'ops_task_lineage', 'ops_problem_case', 'ops_analysis_path',
            'ops_root_cause_type', 'ops_root_cause_analysis', 'ops_analysis_trace_log',
            'ops_case_usage_stats', 'ops_user_feedback',
        ]),
        ('月账管理', [
            'ma_month_account_config', 'ma_task_monitor', 'ma_audit_result',
            'ma_adjustment_record', 'ma_daily_report', 'ma_summary_report',
            'ma_kpi_metrics', 'ma_milestone_track', 'ma_alert_record', 'ma_ml_model_config',
        ]),
        ('系统管理', [
            'sys_user', 'sys_role', 'sys_user_role', 'sys_permission', 'sys_role_permission',
            'sys_department', 'sys_audit_log', 'sys_config', 'sys_data_dict', 'sys_notification_record',
        ]),
        ('AI 助手', [
            'ai_chat_session', 'ai_chat_message', 'sys_ai_config', 'sys_interface_log', 'sys_job_schedule',
        ]),
    ]
    for domain, tbls in tables:
        add_paragraph_zh(doc, domain + '：')
        for t in tbls:
            add_paragraph_zh(doc, '  • ' + t)
    add_heading_zh(doc, '4.2 ER 关系说明', level=2)
    add_paragraph_zh(doc, '核心关系：')
    add_paragraph_zh(doc, '• dq_audit_field_config ←→ dq_audit_rule_config（1:N）')
    add_paragraph_zh(doc, '• dq_audit_task_config ←→ dq_audit_execution（1:N）')
    add_paragraph_zh(doc, '• dq_audit_execution ←→ dq_audit_exception（1:N）')
    add_paragraph_zh(doc, '• sys_user ←→ sys_role（N:M，通过 sys_user_role）')
    add_paragraph_zh(doc, '• sys_role ←→ sys_permission（N:M，通过 sys_role_permission）')
    add_paragraph_zh(doc, '• ops_problem_case ←→ ops_root_cause_analysis（1:N）')
    add_paragraph_zh(doc, '• ai_chat_session ←→ ai_chat_message（1:N）')
    add_paragraph_zh(doc, '完整建表语句见独立文件：db_schema.sql')
    doc.add_page_break()

    # Chapter 5
    add_heading_zh(doc, '5 接口设计概览', level=1)
    add_paragraph_zh(doc, '后端采用 RESTful API 风格，基于 FastAPI 框架，接口统一返回格式：')
    add_paragraph_zh(doc, '{"code": 0, "message": "success", "data": { ... } }')
    add_heading_zh(doc, '5.1 主要接口分组', level=2)
    apis = [
        ('/audit/*', '数据质量稽核接口：字段配置、规则管理、任务执行、告警管理、报告生成'),
        ('/root-cause/*', '根因分析接口：知识库、案例分析、智能分析、改进建议'),
        ('/monthly/*', '月账管理接口：进度监控、日报、报表、月账进度、月账配置'),
        ('/config/*', '配置接口：阶段、里程碑、作业计划、任务、模板解析'),
        ('/settings/*', '系统设置接口：用户、角色、权限、工具、Prompt、通知'),
        ('/ai/*', 'AI 助手接口：会话、消息、配置'),
    ]
    for path, desc in apis:
        add_paragraph_zh(doc, path + '：' + desc)
    add_heading_zh(doc, '5.2 认证方式', level=2)
    add_paragraph_zh(doc, '采用 JWT Token 认证，请求头携带：Authorization: Bearer <token>')
    add_paragraph_zh(doc, 'Token 有效期：24 小时')

    out_path = os.path.join(os.path.dirname(__file__), '产品详细设计文档.docx')
    doc.save(out_path)
    print(f'Word document saved: {out_path}')

if __name__ == '__main__':
    main()
